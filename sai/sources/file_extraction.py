import logging
import time
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from config.errors import DomainError

logger = logging.getLogger(__name__)

LOCAL_DOCUMENT_PART_SIZE = 4000

# 페이지 추출은 페이지마다 걸리는 시간이 크게 다르다. 몇 장째에서 물렸는지 남긴다.
PDF_LOG_EVERY = 50


class FileExtractionError(DomainError):
    pass


def _text_file(data):
    try:
        return data.decode('utf-8-sig')
    except UnicodeDecodeError as exc:
        raise FileExtractionError('unsupported_encoding') from exc


def _pdf_file(data):
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise FileExtractionError('encrypted_file')

        total = len(reader.pages)
        logger.info('PDF 페이지 추출 시작 pages=%d', total)
        parts = []
        for index, page in enumerate(reader.pages, start=1):
            parts.append(page.extract_text() or '')
            if index % PDF_LOG_EVERY == 0:
                logger.info('PDF 페이지 추출 중 %d/%d', index, total)

        return '\n\n'.join(parts)
    except FileExtractionError:
        raise
    except Exception as exc:
        raise FileExtractionError('file_read_failed') from exc


def _docx_file(data):
    try:
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]

        # 사내 규칙이 표로 정리된 문서도 빠뜨리지 않는다.
        for table in document.tables:
            for row in table.rows:
                parts.append(' | '.join(cell.text for cell in row.cells))

        return '\n'.join(parts)
    except Exception as exc:
        raise FileExtractionError('file_read_failed') from exc


def extract_file_text(file_name, data):
    extension = Path(file_name).suffix.lower()

    # 파서가 물리면 여기서 CPU 를 물고 놓지 않는다. pypdf 와 python-docx 는
    # 손상된 파일에서 아주 오래 돌 수 있고, except 로는 그것을 끊지 못한다.
    # 시작 줄만 남고 완료 줄이 없으면 그 파일이 범인이다.
    logger.info(
        '파일 추출 시작 name=%s ext=%s bytes=%d', file_name, extension, len(data)
    )
    started = time.monotonic()

    if extension in {'.txt', '.md'}:
        text = _text_file(data)
    elif extension == '.pdf':
        text = _pdf_file(data)
    elif extension == '.docx':
        text = _docx_file(data)
    else:
        raise FileExtractionError('file_type_not_supported')

    logger.info(
        '파일 추출 완료 name=%s 소요=%.1fs chars=%d',
        file_name, time.monotonic() - started, len(text),
    )

    text = text.strip()
    if not text:
        raise FileExtractionError('empty_document')

    return text


def split_file_text(text, part_size=LOCAL_DOCUMENT_PART_SIZE):
    parts = []
    current = ''

    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue

        line_parts = [
            line[start:start + part_size]
            for start in range(0, len(line), part_size)
        ]
        for line_part in line_parts:
            candidate = f'{current}\n{line_part}'.strip() if current else line_part
            if len(candidate) <= part_size:
                current = candidate
                continue

            parts.append(current)
            current = line_part

    if current:
        parts.append(current)

    return parts
